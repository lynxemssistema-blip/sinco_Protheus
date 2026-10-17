import openpyxl, glob
from collections import defaultdict

files = glob.glob(r'C:\Users\lynx\Downloads\Rela*.xlsx')
f = files[1]
wb = openpyxl.load_workbook(f, data_only=True)
ws = wb[wb.sheetnames[0]]

grupos = defaultdict(lambda: {'desc': '', 'linhas': [], 'qtde_total': 0, 'custo_total': 0.0, 'ocs': set(), 'fornecedores': set(), 'situacoes': set()})

for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0: continue
    if len(row) < 12: continue
    oc, item, produto, desc, fornecedor, dt_em, dt_neg, resp, qtde, custo_uni, custo_total, situacao = row[:12]
    arm = row[12] if len(row) > 12 else None
    
    if produto:
        g = grupos[produto]
        if not g['desc']: g['desc'] = desc
        g['qtde_total'] += (qtde or 0)
        g['custo_total'] += (custo_total or 0)
        if oc: g['ocs'].add(oc)
        if fornecedor: g['fornecedores'].add(fornecedor)
        if situacao: g['situacoes'].add(situacao)
        g['linhas'].append({'oc': oc, 'item': item, 'qtde': qtde, 'custo_uni': custo_uni, 'custo_total': custo_total, 'situacao': situacao, 'dt_neg': dt_neg})

total_qtde = sum(g['qtde_total'] for g in grupos.values())
total_custo = sum(g['custo_total'] for g in grupos.values())
print("Total grupos (produtos unicos): " + str(len(grupos)))
print("Total QTDE_EM_ABERTO geral: " + str(total_qtde))
print("Total CUSTO_TOTAL geral: R$" + "{:,.2f}".format(total_custo))

print('\nTop 5 grupos por custo:')
for prod, g in sorted(grupos.items(), key=lambda x: x[1]['custo_total'], reverse=True)[:5]:
    print("  " + str(prod) + " | " + str(g['desc'])[:40] + " | qtde:" + str(g['qtde_total']) + " | custo:R$" + "{:,.2f}".format(g['custo_total']) + " | OCs:" + str(list(g['ocs'])))

print('\nAnalise prefixos (2 primeiros segmentos = familia de produto):')
prefix_groups = defaultdict(lambda: {'count': 0, 'qtde': 0, 'custo': 0.0})
for prod, g in grupos.items():
    parts = str(prod).split('.')
    prefix = '.'.join(parts[:2]) if len(parts) >= 2 else prod
    prefix_groups[prefix]['count'] += 1
    prefix_groups[prefix]['qtde'] += g['qtde_total']
    prefix_groups[prefix]['custo'] += g['custo_total']

print("Total prefixos unicos (familias): " + str(len(prefix_groups)))
for pref, pg in sorted(prefix_groups.items(), key=lambda x: x[1]['custo'], reverse=True)[:10]:
    print("  " + str(pref) + " | produtos:" + str(pg['count']) + " | qtde:" + str(pg['qtde']) + " | custo:R$" + "{:,.2f}".format(pg['custo']))
