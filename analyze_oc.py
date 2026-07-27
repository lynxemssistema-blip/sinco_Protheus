import openpyxl, glob
from collections import defaultdict

files = glob.glob(r'C:\Users\lynx\Downloads\Rela*.xlsx')
f = files[1]
print(f"Opening: {f}")
wb = openpyxl.load_workbook(f, data_only=True)
sheet_name = wb.sheetnames[0]
ws = wb[sheet_name]

# Analyze col C (PRODUTO) patterns for grouping
print('=== Coluna C (PRODUTO) - análise de padrões ===')
produtos = defaultdict(lambda: {'count': 0, 'qtde_total': 0, 'custo_total': 0.0, 'exemplos': []})
fornecedores = set()
situacoes = set()
responsaveis = set()
ocs_unicas = set()

for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0: continue  # skip header
    if len(row) < 12: continue
    oc, item, produto, desc, fornecedor, dt_em, dt_neg, resp, qtde, custo_uni, custo_total, situacao = row[:12]
    arm = row[12] if len(row) > 12 else None
    
    if oc: ocs_unicas.add(oc)
    if produto:
        produtos[produto]['count'] += 1
        produtos[produto]['qtde_total'] += (qtde or 0)
        produtos[produto]['custo_total'] += (custo_total or 0)
        if len(produtos[produto]['exemplos']) < 2:
            produtos[produto]['exemplos'].append({'desc': desc, 'fornecedor': fornecedor, 'oc': oc})
    if fornecedor: fornecedores.add(fornecedor)
    if situacao: situacoes.add(situacao)
    if resp: responsaveis.add(resp)

print(f'Total linhas (excl. header): {ws.max_row - 1}')
print(f'Total OCs únicas: {len(ocs_unicas)}')
print(f'Total produtos únicos (Col C): {len(produtos)}')
print(f'Fornecedores únicos: {len(fornecedores)}')
print(f'Situações: {situacoes}')
print(f'Responsáveis: {responsaveis}')

print('\n=== Amostra de OCs únicas ===')
for oc in list(sorted(ocs_unicas))[:10]:
    print(f'  OC: {oc}')

print('\n=== Top 15 produtos por quantidade em aberto (Col I = QTDE_EM_ABERTO) ===')
top_prod = sorted(produtos.items(), key=lambda x: x[1]['qtde_total'], reverse=True)[:15]
for prod, data in top_prod:
    print(f'  {prod} | qtde: {data["qtde_total"]:,} | custo_total: R${data["custo_total"]:,.2f} | linhas: {data["count"]}')
    if data['exemplos']:
        print(f'    ex: {data["exemplos"][0]["desc"][:60]}')

print('\n=== Amostra de 20 códigos PRODUTO (Col C) ===')
for i, (prod, data) in enumerate(list(produtos.items())[:20]):
    print(f'  [{i+1}] {prod} -> {data["exemplos"][0]["desc"][:50] if data["exemplos"] else "-"}')

print('\n=== Formato dos códigos de produto (analisando padrão) ===')
# Check for VMI-style codes
vmi_codes = [p for p in produtos.keys() if '.' in str(p) or '-' in str(p)]
non_vmi = [p for p in produtos.keys() if '.' not in str(p) and '-' not in str(p)]
print(f'Códigos com "." ou "-": {len(vmi_codes)} exemplos: {vmi_codes[:5]}')
print(f'Códigos simples: {len(non_vmi)} exemplos: {non_vmi[:5]}')
