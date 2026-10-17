import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_irricom_presentation(output_path):
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    
    # Paleta de Cores Suaves e Leves (Sinco Executive Light)
    BG_COLOR = RGBColor(248, 250, 252)        # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)         # #FFFFFF
    CARD_BORDER = RGBColor(226, 232, 240)     # #E2E8F0
    TEXT_MAIN = RGBColor(30, 41, 37)          # #1E2925 (Verde Ardósia Escuro)
    TEXT_MUTED = RGBColor(100, 116, 139)      # #64748B (Slate médio)
    TEXT_LIGHT = RGBColor(148, 163, 184)      # #94A3B8
    BRAND_FOREST = RGBColor(37, 68, 56)       # #254438 (Verde Floresta institucional)
    BRAND_SAGE = RGBColor(16, 149, 106)       # #10956A (Verde Sálvia)
    BRAND_LIME = RGBColor(184, 194, 10)       # #B8C20A (Lima suave institucional)
    TINT_GREEN_BG = RGBColor(236, 253, 245)   # #ECFDF5 (Fundo leve esmeralda)
    TINT_AMBER_BG = RGBColor(254, 243, 199)   # #FEF3C7 (Dourado suave)
    TEXT_AMBER = RGBColor(180, 83, 9)         # #B45309
    
    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, category_text, title_text, subtitle_text=""):
        # Barra superior decorativa
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(11.733), Inches(0.85))
        bar.fill.solid()
        bar.fill.fore_color.rgb = CARD_BG
        bar.line.color.rgb = CARD_BORDER
        bar.line.width = Pt(1)
        
        # Pill Tag da Categoria
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.68), Inches(2.2), Inches(0.35))
        pill.fill.solid()
        pill.fill.fore_color.rgb = TINT_GREEN_BG
        pill.line.color.rgb = BRAND_SAGE
        pill.line.width = Pt(1)
        tf_pill = pill.text_frame
        tf_pill.word_wrap = False
        p_pill = tf_pill.paragraphs[0]
        p_pill.text = category_text.upper()
        p_pill.font.size = Pt(9.5)
        p_pill.font.bold = True
        p_pill.font.color.rgb = BRAND_SAGE
        p_pill.alignment = PP_ALIGN.CENTER
        
        # Título do Header
        tx_box = slide.shapes.add_textbox(Inches(3.35), Inches(0.55), Inches(6.8), Inches(0.45))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        
        if subtitle_text:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(9.5)
            p_sub.font.color.rgb = TEXT_MUTED
            
        # Selo Lynx EMS / Sinco à direita
        lynx_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), Inches(0.65), Inches(1.8), Inches(0.4))
        lynx_badge.fill.solid()
        lynx_badge.fill.fore_color.rgb = BRAND_FOREST
        lynx_badge.line.fill.background()
        tf_lynx = lynx_badge.text_frame
        p_lynx = tf_lynx.paragraphs[0]
        p_lynx.text = "LYNX EMS • SINCO"
        p_lynx.font.size = Pt(9)
        p_lynx.font.bold = True
        p_lynx.font.color.rgb = RGBColor(249, 248, 243)
        p_lynx.alignment = PP_ALIGN.CENTER

    def add_footer(slide, current_page, total_pages=12):
        # Linha fina sutil
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.01))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()
        
        # Textos de rodapé
        tb_left = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(6.0), Inches(0.35))
        p_l = tb_left.text_frame.paragraphs[0]
        p_l.text = "Proposta Técnico-Comercial • Implementação Ecossistema SINCO • Irricom"
        p_l.font.size = Pt(9)
        p_l.font.color.rgb = TEXT_LIGHT
        
        tb_right = slide.shapes.add_textbox(Inches(9.5), Inches(6.95), Inches(3.0), Inches(0.35))
        p_r = tb_right.text_frame.paragraphs[0]
        p_r.text = f"Página {current_page} de {total_pages} | Confidencial"
        p_r.font.size = Pt(9)
        p_r.font.color.rgb = TEXT_LIGHT
        p_r.alignment = PP_ALIGN.RIGHT

    def create_card(slide, left, top, width, height, title, items, tag="", border_color=CARD_BORDER, bg_color=CARD_BG):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)
        
        if tag:
            p_tag = tf.paragraphs[0]
            p_tag.text = tag.upper()
            p_tag.font.size = Pt(8.5)
            p_tag.font.bold = True
            p_tag.font.color.rgb = BRAND_SAGE
            p_tag.space_after = Pt(4)
            p_title = tf.add_paragraph()
        else:
            p_title = tf.paragraphs[0]
            
        p_title.text = title
        p_title.font.size = Pt(13)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_MAIN
        p_title.space_after = Pt(8)
        
        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = f"•  {item}"
            p_item.font.size = Pt(10)
            p_item.font.color.rgb = TEXT_MUTED
            p_item.space_after = Pt(4)
            
        return card

    # =========================================================================
    # SLIDE 1: CAPA EXECUTIVA
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    
    cover_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.333), Inches(5.5))
    cover_box.fill.solid()
    cover_box.fill.fore_color.rgb = CARD_BG
    cover_box.line.color.rgb = CARD_BORDER
    cover_box.line.width = Pt(1.5)
    
    tag_top = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.2), Inches(1.5), Inches(4.5), Inches(0.4))
    tag_top.fill.solid()
    tag_top.fill.fore_color.rgb = TINT_GREEN_BG
    tag_top.line.color.rgb = BRAND_SAGE
    tf_tag = tag_top.text_frame
    p_tt = tf_tag.paragraphs[0]
    p_tt.text = "PROPOSTA TÉCNICO-COMERCIAL • 2026"
    p_tt.font.size = Pt(10)
    p_tt.font.bold = True
    p_tt.font.color.rgb = BRAND_SAGE
    p_tt.alignment = PP_ALIGN.CENTER
    
    tb_cov = s1.shapes.add_textbox(Inches(2.2), Inches(2.1), Inches(8.8), Inches(3.2))
    tf_c = tb_cov.text_frame
    tf_c.word_wrap = True
    
    p1 = tf_c.paragraphs[0]
    p1.text = "Implementação do Ecossistema\nSINCO"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = BRAND_FOREST
    p1.space_after = Pt(12)
    
    p2 = tf_c.add_paragraph()
    p2.text = "Solução Sob Medida para Gestão Integrada de Engenharia, Fabricação, SST, SGQ e Mobilidade de Campo."
    p2.font.size = Pt(13)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(24)
    
    p3 = tf_c.add_paragraph()
    p3.text = "Cliente Alvo: IRRICOM (Indústria, Mineração & Irrigação)\nElaborado por: LYNX Soluções Corporativas • Sinco Industrial Suite"
    p3.font.size = Pt(11)
    p3.font.bold = True
    p3.font.color.rgb = TEXT_MAIN

    tb_fc = s1.shapes.add_textbox(Inches(2.2), Inches(5.8), Inches(8.8), Inches(0.4))
    p_fc = tb_fc.text_frame.paragraphs[0]
    p_fc.text = "Ambiente Dedicado: https://irricom.lynxems.com.br • Metodologia Ágil (MVP) • Sustentação SaaS"
    p_fc.font.size = Pt(9.5)
    p_fc.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 2: CONTEXTO DA IRRICOM & ALINHAMENTO
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Contexto Corporativo", "A Operação da Irricom & Alinhamento Estratégico", "30 anos de excelência conectando engenharia, mineração, montagem eletromecânica e irrigação")
    
    create_card(s2, 0.8, 1.6, 3.65, 4.9, "Escala Operacional", [
        "30+ anos de história no mercado nacional.",
        "+4.000 projetos executados com sucesso.",
        "+640 colaboradores e equipes em campo.",
        "Presença nas maiores mineradoras do Brasil.",
        "Distribuição autorizada Premium Rain Bird."
    ], tag="Portfólio & História")

    create_card(s2, 4.8, 1.6, 3.65, 4.9, "Desafio de Governança", [
        "Certificação NBR ISO 9001 há 8+ anos.",
        "Exigência rigorosa de SST em mineração.",
        "Comunicação crítica: Matriz vs. Fábrica vs. Canteiros.",
        "Falta de ferramentas táteis para a ponta.",
        "Eliminação definitiva de papel e planilhas."
    ], tag="Qualidade & SST")

    create_card(s2, 8.8, 1.6, 3.65, 4.9, "Objetivo da Proposta", [
        "Implantação sob medida do Ecossistema SINCO.",
        "Mobilidade real para apontamentos e requisições.",
        "Rastreabilidade de ponta a ponta na nuvem.",
        "Setup ágil com desenvolvimento imersivo.",
        "Infraestrutura 100% gerenciada pela Lynx."
    ], tag="Parceria Lynx + Irricom", border_color=BRAND_SAGE)
    
    add_footer(s2, 2)

    # =========================================================================
    # SLIDE 3: O ECOSSISTEMA SINCO INDUSTRIAL SUITE
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Tecnologia Sinco", "O Que é o Ecossistema SINCO", "A inteligência operacional compatível com qualquer ERP (Protheus, Omie, Sankhya, etc.)")

    create_card(s3, 0.8, 1.6, 3.65, 4.9, "Mobilidade Operacional", [
        "Interfaces táteis intuitivas para tablet e smartphone.",
        "Apontamentos diretos no chão de fábrica e na obra.",
        "Eliminação de atrasos de digitação e retrabalho.",
        "Operação simplificada para operadores e encarregados."
    ], tag="Chão de Fábrica & Obra")

    create_card(s3, 4.8, 1.6, 3.65, 4.9, "Integração Protheus", [
        "Conexão nativa e bidirecional com o Protheus.",
        "Roteiros, OPs, produtos e clientes sincronizados.",
        "Rastreabilidade de insumos e matérias-primas.",
        "Sem necessidade de customizações invasivas no ERP."
    ], tag="Sinergia com ERP")

    create_card(s3, 8.8, 1.6, 3.65, 4.9, "Cockpit Gerencial", [
        "Dashboards em tempo real para matriz e diretoria.",
        "Visibilidade de gargalos, avanço físico e pendências.",
        "Conformidade com normas de governança e auditoria.",
        "Histórico digital auditável e segurança avançada."
    ], tag="Controle & Gestão", border_color=BRAND_SAGE)

    add_footer(s3, 3)

    # =========================================================================
    # SLIDE 4: ESCOPO FASE 1 - SST E SGQ
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Fase 1 • Escopo", "Módulos de SST e SGQ (Qualidade & Segurança)", "Garantindo compliance total em áreas industriais e atendimento à ISO 9001")

    create_card(s4, 0.8, 1.6, 5.7, 4.9, "Módulo SST (Saúde e Segurança)", [
        "Controle rigoroso para atendimento a NRs em áreas de mineração.",
        "Registro digital de inspeções, DDS e checagens de segurança.",
        "Gestão de conformidade de EPIs e autorizações de trabalho.",
        "Alerta preventivo de incidentes e bloqueios de segurança.",
        "Proteção integral para os mais de 640 colaboradores da Irricom."
    ], tag="Segurança do Trabalho", border_color=BRAND_SAGE)

    create_card(s4, 6.8, 1.6, 5.7, 4.9, "Módulo SGQ (Gestão da Qualidade)", [
        "Suporte digital estruturado às exigências da NBR ISO 9001.",
        "Gestão digital de não conformidades (RNCs) e planos de ação.",
        "Controle de documentos, versões de engenharia e procedimentos.",
        "Mapeamento de processos e calibração de instrumentos de precisão.",
        "Prontidão contínua para auditorias externas sem risco de apontamentos."
    ], tag="Qualidade ISO 9001", border_color=BRAND_FOREST)

    add_footer(s4, 4)

    # =========================================================================
    # SLIDE 5: ESCOPO FASE 1 - FÁBRICA, CAMPO E SOLICITAÇÕES
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Fase 1 • Escopo", "Gestão Fabril, Gestão de Campo & Solicitações", "Da manufatura interna ao avanço físico em canteiros de obra remotos")

    create_card(s5, 0.8, 1.6, 3.65, 4.9, "Gestão Fabril", [
        "Controle da linha de produção e montagem interna.",
        "Apontamento de ordens de serviço e etapas.",
        "Rastreabilidade de equipamentos customizados.",
        "Medição de tempos-padrão e produtividade fabril."
    ], tag="Manufatura & Oficina")

    create_card(s5, 4.8, 1.6, 3.65, 4.9, "Gestão de Campo", [
        "Acompanhamento do avanço físico de obras e montagens.",
        "Diário de obra digital com fotos e apontamentos.",
        "Gestão da presença e produtividade de equipes de campo.",
        "Visibilidade em tempo real para os gestores na matriz."
    ], tag="Obras & Montagens")

    create_card(s5, 8.8, 1.6, 3.65, 4.9, "Solicitações de Campo", [
        "Canal digital rápido e oficial Canteiro ↔ Matriz.",
        "Requisição de materiais, peças de reposição e ferramentas.",
        "Eliminação de pedidos perdidos em WhatsApp pessoal.",
        "Histórico e rastreamento do status de atendimento."
    ], tag="Comunicação Oficial", border_color=BRAND_SAGE)

    add_footer(s5, 5)

    # =========================================================================
    # SLIDE 6: FASE 2 - ROADMAP DE EXPANSÃO (FORA DO ESCOPO ATUAL)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Fase 2 • Futuro", "Roadmap de Expansão (Módulos Futuros)", "Evolução planejada do ecossistema após consolidação do MVP")

    create_card(s6, 0.8, 1.6, 3.65, 4.9, "Módulo Almoxarifado", [
        "Controle de saldos físicos e inventários periódicos.",
        "Rastreabilidade de caixas de ferramentas e ativos em campo.",
        "Entradas, saídas e transferências entre obras.",
        "Previsão de ativação em etapas futuras."
    ], tag="Roadmap Futuro")

    create_card(s6, 4.8, 1.6, 3.65, 4.9, "Módulo Compras", [
        "Fluxo digital de cotações e pedidos de compra.",
        "Auditoria de pedidos de compra versus XML de fornecedores.",
        "Gestão de prazos de entrega e fornecedores homologados.",
        "Previsão de ativação em etapas futuras."
    ], tag="Roadmap Futuro")

    create_card(s6, 8.8, 1.6, 3.65, 4.9, "Gestão Financeira", [
        "Fluxo de caixa gerencial e controle de pagamentos.",
        "Controle de faturamento atrelado a medições de obra.",
        "Indicadores financeiros e DRE gerencial.",
        "Previsão de ativação em etapas futuras."
    ], tag="Roadmap Futuro")

    add_footer(s6, 6)

    # =========================================================================
    # SLIDE 7: METODOLOGIA DE IMPLANTAÇÃO - INOVAÇÃO ÁGIL (MVP)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Metodologia Ágil", "A Metodologia LYNX: Inovação Ágil baseada em MVP", "Desenvolvimento profundo e orgânico, sem modelos tradicionais engessados")

    create_card(s7, 0.8, 1.6, 3.65, 4.9, "Imersão de 60 Dias", [
        "Dedicação intensiva e exclusiva da Lynx nos primeiros 2 meses.",
        "Adaptação contínua às rotinas reais e peculiaridades da Irricom.",
        "Engenharia de software focada nos gargalos operacionais.",
        "Sem limites burocráticos de escopo fechado nesta fase."
    ], tag="Dedicação Profunda", border_color=BRAND_SAGE)

    create_card(s7, 4.8, 1.6, 3.65, 4.9, "Aceleração de Valor", [
        "O sistema NÃO precisa estar 100% finalizado para começar a rodar.",
        "Liberação em ciclos curtos para uso imediato das equipes.",
        "Feedback prático vindo diretamente do campo e da fábrica.",
        "Ajustes rápidos de layout e regras de negócio em tempo real."
    ], tag="Uso Imediato")

    create_card(s7, 8.8, 1.6, 3.65, 4.9, "Foco na Assertividade", [
        "A realidade no campo muda: o software acompanha a obra.",
        "Garantia de que a ferramenta resolva dores reais sem atrito.",
        "Adesão natural dos colaboradores através de usabilidade simples.",
        "Implantação assertiva e com retorno sobre o investimento imediato."
    ], tag="Eficiência Real")

    add_footer(s7, 7)

    # =========================================================================
    # SLIDE 8: CHECKPOINT DE AVALIAÇÃO & CONTINUIDADE
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Governança do Projeto", "Checkpoint de Avaliação aos 60 Dias", "Transparência total e validação prática de aderência antes da fase continuada")

    create_card(s8, 0.8, 1.6, 5.7, 4.9, "Relatório Técnico de Entrega", [
        "Ao término dos 60 dias de imersão, a Lynx emitirá um dossiê completo.",
        "Detalhamento do status de cada funcionalidade entregue e operante.",
        "Métricas de utilização nos canteiros e na fábrica.",
        "Mapeamento de ganhos de tempo e eliminação de processos em papel.",
        "Plano de sustentação e evolução para os meses seguintes."
    ], tag="Documentação Formal", border_color=BRAND_FOREST)

    create_card(s8, 6.8, 1.6, 5.7, 4.9, "Pesquisa de Satisfação & Validação", [
        "Coleta formal de feedback com os líderes e usuários-chave da Irricom.",
        "Avaliação prática da aderência da plataforma à rotina das obras.",
        "Decisão segura e fundamentada para continuidade do ecossistema.",
        "Garantia de parceria de longo prazo baseada em resultados comprovados."
    ], tag="Validação Irricom", border_color=BRAND_SAGE)

    add_footer(s8, 8)

    # =========================================================================
    # SLIDE 9: INFRAESTRUTURA TECNOLÓGICA INCLUSA (SAAS)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Infraestrutura SaaS", "Infraestrutura Tecnológica Gerenciada pela Lynx", "A Irricom foca na engenharia de ponta; a Lynx cuida de toda a sustentação técnica")

    create_card(s9, 0.8, 1.6, 2.75, 4.9, "Domínio Exclusivo", [
        "Acesso dedicado via:",
        "https://irricom.lynxems.com.br",
        "Ambiente web moderno, rápido e seguro.",
        "Criptografia SSL ponta a ponta."
    ], tag="Acesso Web")

    create_card(s9, 3.8, 1.6, 2.75, 4.9, "Servidores & Nuvem", [
        "Servidor web e banco de dados em nuvem de alta disponibilidade.",
        "Monitoramento contínuo de uptime e desempenho.",
        "Sem necessidade de servidores locais na Irricom."
    ], tag="Cloud Hosting")

    create_card(s9, 6.8, 1.6, 2.75, 4.9, "Backups Automáticos", [
        "Rotinas automáticas de backup de segurança diário.",
        "Políticas de recuperação de dados e retenção segura.",
        "Proteção dos registros operacionais e relatórios."
    ], tag="Segurança de Dados")

    create_card(s9, 9.8, 1.6, 2.75, 4.9, "Manutenção & Evolução", [
        "Todo o custo de desenvolvimento contínuo incluso.",
        "Correções preventivas e melhorias de usabilidade.",
        "Suporte técnico especializado com a equipe da Lynx."
    ], tag="Suporte Contínuo", border_color=BRAND_SAGE)

    add_footer(s9, 9)

    # =========================================================================
    # SLIDE 10: CONDIÇÕES COMERCIAIS - SETUP INICIAL & IMERSÃO
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Investimento • Fase 1", "Imersão e Setup Inicial (Fase Ágil de 60 Dias)", "Dedicação intensiva de desenvolvimento e parametrização profunda")

    create_card(s10, 1.5, 1.8, 4.8, 4.6, "Mês 1 • Setup & Imersão", [
        "Valor: R$ 16.520,60 (Impostos inclusos)",
        "Atividades abrangidas:",
        "  - Levantamento detalhado de processos com as equipes da Irricom.",
        "  - Configuração do ambiente dedicado irricom.lynxems.com.br.",
        "  - Implantação e customização dos módulos SST e SGQ.",
        "  - Início dos testes práticos com equipes-piloto."
    ], tag="Primeiros 30 Dias", border_color=BRAND_SAGE)

    create_card(s10, 7.0, 1.8, 4.8, 4.6, "Mês 2 • Expansão & Validação", [
        "Valor: R$ 16.520,60 (Impostos inclusos)",
        "Atividades abrangidas:",
        "  - Desenvolvimento e liberação de Gestão Fabril e Gestão de Campo.",
        "  - Canal de Solicitações de Obra operando entre canteiros e matriz.",
        "  - Ajustes baseados no feedback do dia a dia da obra.",
        "  - Emissão do Relatório Técnico de Entrega e Checkpoint."
    ], tag="31 a 60 Dias", border_color=BRAND_FOREST)

    add_footer(s10, 10)

    # =========================================================================
    # SLIDE 11: SUSTENTAÇÃO & REGRA DE ESCALABILIDADE DE CANTEIROS
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Investimento • Fase 2", "Sustentação, Licenciamento & Regra de Canteiros", "A partir do 3º mês: modelo flexível que acompanha organicamente o volume da Irricom")

    create_card(s11, 0.8, 1.6, 5.7, 4.9, "Sustentação Base (Franquia de 15 Canteiros)", [
        "Vigor: A partir do 3º mês (após aprovação do checkpoint).",
        "Valor Fixo Base: Equivalente a 6 (seis) Salários Mínimos vigentes.",
        "Franquia Inclusa: Até 15 canteiros de obras cadastrados e ativos simultaneamente.",
        "Abrange: Infraestrutura em nuvem, backups, licenças ilimitadas de usuários e melhorias contínuas do sistema."
    ], tag="Mensalidade Base", border_color=BRAND_FOREST)

    create_card(s11, 6.8, 1.6, 5.7, 4.9, "Regra de Escalabilidade (Pay-per-use)", [
        "Acréscimo: +0,5 (meio) Salário Mínimo por canteiro extra que exceder os 15 da franquia.",
        "Autonomia e Transparência Total: A Irricom tem controle exclusivo de abertura e encerramento de canteiros no sistema.",
        "Cessação Imediata: Ao inativar ou encerrar a obra no app, a cobrança adicional cessa imediatamente no mês seguinte.",
        "Justiça Comercial: A Irricom paga apenas pelo volume real de obras que estiver executando."
    ], tag="Pay-Per-Use Justo", border_color=BRAND_SAGE)

    add_footer(s11, 11)

    # =========================================================================
    # SLIDE 12: FATURAMENTO & PRÓXIMOS PASSOS
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Condições Comerciais", "Regras de Faturamento e Próximos Passos", "Formalização transparente para início imediato do cronograma de imersão")

    create_card(s12, 0.8, 1.6, 5.7, 4.9, "Regras de Faturamento", [
        "Emissão da Nota Fiscal: Todo dia 28 de cada mês.",
        "Vencimento do Pagamento: Previsto para o dia 10 do mês subsequente.",
        "Tributos: Todos os impostos aplicáveis já estão devidamente contemplados.",
        "Canal Financeiro Oficial: Envio eletrônico com boleto bancário e chave PIX corporativa."
    ], tag="Faturamento & Prazos")

    create_card(s12, 6.8, 1.6, 5.7, 4.9, "Cronograma de Kick-Off", [
        "Passo 1: Aprovação e formalização da proposta comercial.",
        "Passo 2: Provisionamento do ambiente irricom.lynxems.com.br.",
        "Passo 3: Reunião de alinhamento com a liderança da Irricom.",
        "Passo 4: Início da imersão de 60 dias e sprints ágeis.",
        "Passo 5: Liberação dos primeiros módulos operacionais em campo."
    ], tag="Início dos Trabalhos", border_color=BRAND_SAGE)

    add_footer(s12, 12)

    # Salva a apresentação
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "client/public/propostas/Proposta_Comercial_SINCO_Irricom_2026.pptx"
    create_irricom_presentation(out)
