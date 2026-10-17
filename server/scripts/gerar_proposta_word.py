import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import json
import os
import sys
from datetime import datetime, timedelta

def set_cell_background(cell, hex_color):
    """Define a cor de fundo de uma célula da tabela."""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Define margens internas (padding) da célula em dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    """Define bordas sutis em uma tabela do Word."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideV w:val="none"/>'
                        f'<w:left w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tblBorders>')
    tblPr.append(borders)

def fmt_br(val):
    """Formata número para padrão brasileiro R$."""
    try:
        n = float(val)
        return f"{n:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except:
        return "0,00"

def parse_date_safe(d_str):
    """Tenta converter string de data YYYY-MM-DD em objeto datetime."""
    if not d_str:
        return datetime.now()
    try:
        return datetime.strptime(str(d_str)[:10], "%Y-%m-%d")
    except:
        try:
            return datetime.strptime(str(d_str)[:10], "%d/%m/%Y")
        except:
            return datetime.now()

def generate_proposal_docx(json_path=None, output_path=None):
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    if not json_path:
        json_path = os.path.join(base_dir, '..', 'tmp_proposta.json')
    if not output_path:
        output_path = os.path.join(base_dir, '..', 'Proposta_Comercial_SINCO.docx')

    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # Identificadores e parceiros
    cliente_nome = data.get('cliente_nome', 'Cliente')
    cliente_razao = data.get('cliente_razao_social', cliente_nome)
    codigo = data.get('codigo', 'PROP-2026-001')
    revisao = data.get('revisao', 'Rev 00')
    data_proposta = data.get('data_proposta', datetime.now().strftime('%Y-%m-%d'))
    url_ambiente = data.get('url_ambiente', f"https://{cliente_nome.lower().replace(' ', '')}.lynxems.com.br")

    fornecedor = data.get('fornecedor', {}) or {}
    fornecedor_nome = fornecedor.get('empresa', 'LYNX EMS SISTEMA & CONSULTORIA')

    parceira = data.get('empresa_parceira', {}) or {}
    parceira_habilitada = parceira.get('habilitada', True)
    parceira_nome = parceira.get('nome', 'SIM Consultoria') if parceira_habilitada else ''
    parceira_papel = parceira.get('papel', 'Consultoria Especializada TOTVS Protheus & Implantação de Processos')

    doc = docx.Document()

    # Definir margens do documento (2 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Cores corporativas Lynx & SINCO (Floresta / Sage / Navy)
    COLOR_PRIMARY = RGBColor(14, 43, 92)     # Azul Marinho Corporativo Lynx
    COLOR_FOREST = RGBColor(10, 61, 46)     # Verde Floresta Corporativo SINCO
    COLOR_SAGE = RGBColor(45, 106, 79)      # Verde Sage
    COLOR_TEXT = RGBColor(30, 41, 59)        # Slate 800
    COLOR_MUTED = RGBColor(100, 116, 139)    # Slate 500
    COLOR_AMBER = RGBColor(180, 83, 9)       # Amber 700

    HEX_PRIMARY = "0E2B5C"
    HEX_FOREST = "0A3D2E"
    HEX_LIGHT_BG = "F8FAFC"
    HEX_ACCENT_BG = "ECFDF5"
    HEX_BORDER = "E2E8F0"

    # ==================== CABEÇALHO COM LOGOS LYNX & SIM ====================
    # Tabela de 1 linha e 2 colunas para colocar os logos e metadados
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False

    cell_left = header_table.cell(0, 0)
    cell_right = header_table.cell(0, 1)
    cell_left.width = Inches(4.2)
    cell_right.width = Inches(2.6)
    set_cell_margins(cell_left, 40, 40, 40, 40)
    set_cell_margins(cell_right, 40, 40, 40, 40)

    # Lado esquerdo: Logos da Lynx e da SIM Consultoria (dinâmicos / selecionados)
    p_logos = cell_left.paragraphs[0]
    p_logos.alignment = WD_ALIGN_PARAGRAPH.LEFT

    import base64
    import io

    def resolve_image(doc_run, img_source, default_name, width_inches=1.4):
        sources = [img_source, default_name]
        for src in sources:
            if not src:
                continue
            try:
                if str(src).startswith('data:image'):
                    # Decodifica Base64 enviado pelo navegador
                    header, data_part = src.split(',', 1)
                    img_bytes = base64.b64decode(data_part)
                    stream = io.BytesIO(img_bytes)
                    doc_run.add_picture(stream, width=Inches(width_inches))
                    return True
                else:
                    clean_path = str(src).lstrip('/')
                    candidates = [
                        os.path.join(base_dir, 'public', clean_path),
                        os.path.join(base_dir, '..', 'client', 'public', clean_path),
                        os.path.join(base_dir, 'public', 'logos', os.path.basename(clean_path)),
                        os.path.join(base_dir, '..', 'client', 'public', 'logos', os.path.basename(clean_path)),
                        os.path.join(base_dir, 'public', os.path.basename(clean_path)),
                        os.path.join(base_dir, '..', 'client', 'public', os.path.basename(clean_path)),
                        src
                    ]
                    for cand in candidates:
                        if os.path.exists(cand) and os.path.isfile(cand):
                            doc_run.add_picture(cand, width=Inches(width_inches))
                            return True
            except Exception as e:
                print(f"Aviso ao carregar imagem ({str(src)[:30]}):", e)
        return False

    # Logo Lynx
    lynx_logo_src = fornecedor.get('logo_url') or 'lynx-ems-logo.jpg'
    r_l = p_logos.add_run()
    if resolve_image(r_l, lynx_logo_src, 'lynx-ems-logo.jpg', width_inches=1.35):
        p_logos.add_run("   ")

    # Logo Parceira se habilitada
    if parceira_habilitada:
        parceira_logo_src = parceira.get('logo_url') or 'sim-consultoria-logo.png'
        r_s = p_logos.add_run()
        resolve_image(r_s, parceira_logo_src, 'sim-consultoria-logo.png', width_inches=1.4)

    # Lado direito: Metadados do Documento
    p_meta = cell_right.paragraphs[0]
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_meta.add_run(f"REF: {codigo} • {revisao}\n")
    r1.font.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = COLOR_FOREST

    r2 = p_meta.add_run(f"Data: {data_proposta}\n")
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = COLOR_MUTED

    if parceira_habilitada:
        r3 = p_meta.add_run(f"Consórcio: Lynx EMS & {parceira_nome}\n")
    else:
        r3 = p_meta.add_run("Emissão: Lynx EMS Soluções\n")
    r3.font.size = Pt(8.5)
    r3.font.bold = True
    r3.font.color.rgb = COLOR_PRIMARY

    r4 = p_meta.add_run(f"Ambiente: {url_ambiente}")
    r4.font.size = Pt(8)
    r4.font.color.rgb = COLOR_MUTED

    # Linha divisória
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(8)
    p_div.paragraph_format.space_after = Pt(14)
    r_line = p_div.add_run("―" * 60)
    r_line.font.color.rgb = COLOR_SAGE
    r_line.font.bold = True

    # ==================== TÍTULO E CAPA ====================
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(4)

    run_subt = p_title.add_run("PROPOSTA TÉCNICA E COMERCIAL DE IMPLEMENTAÇÃO\n")
    run_subt.font.size = Pt(11)
    run_subt.font.bold = True
    run_subt.font.color.rgb = COLOR_SAGE

    run_main = p_title.add_run(f"Ecossistema SINCO para a {cliente_nome}\n")
    run_main.font.size = Pt(20)
    run_main.font.bold = True
    run_main.font.color.rgb = COLOR_PRIMARY

    p_dest = doc.add_paragraph()
    p_dest.paragraph_format.space_after = Pt(16)
    r_dest = p_dest.add_run(f"Destinatário: {cliente_razao} | Setor: {data.get('setor', 'Engenharia e Indústria')}")
    r_dest.font.size = Pt(10)
    r_dest.font.color.rgb = COLOR_MUTED

    # Box Informativo de Consórcio / Parceria
    if parceira_habilitada:
        box_table = doc.add_table(rows=1, cols=1)
        box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        box_table.autofit = False
        cell_box = box_table.cell(0, 0)
        cell_box.width = Inches(6.8)
        set_cell_background(cell_box, "ECFDF5")
        set_cell_margins(cell_box, 140, 140, 180, 180)
        p_b = cell_box.paragraphs[0]
        p_b.paragraph_format.space_after = Pt(0)
        r_b_tag = p_b.add_run("★ CONSÓRCIO DE ENTREGA E GOVERNANÇA INTEGRADA\n")
        r_b_tag.font.bold = True
        r_b_tag.font.size = Pt(9.5)
        r_b_tag.font.color.rgb = COLOR_FOREST

        r_b_txt = p_b.add_run(
            f"A presente solução é entregue em parceria estratégica entre a LYNX EMS (responsável pela tecnologia, arquitetura de software SINCO, módulos mobile offline e conectores universais de integração) e a {parceira_nome} ({parceira_papel}), unindo excelência tecnológica à máxima aderência aos processos industriais e conexão com qualquer ERP do mercado (Protheus, Omie, Sankhya, entre outros)."
        )
        r_b_txt.font.size = Pt(9)
        r_b_txt.font.color.rgb = COLOR_TEXT

        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(8)

    # Helpers para seções e parágrafos
    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(title)
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = COLOR_FOREST
        return p

    def add_p(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.bold = True
            r_pre.font.size = Pt(9.5)
            r_pre.font.color.rgb = COLOR_PRIMARY
        r = p.add_run(text)
        r.font.size = Pt(9.5)
        r.font.italic = italic
        r.font.color.rgb = COLOR_TEXT
        return p

    def add_bullet(title, desc):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r_t = p.add_run(title + ": ")
        r_t.font.bold = True
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = COLOR_FOREST
        r_d = p.add_run(desc)
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_TEXT
        return p

    # ==================== SEÇÃO 1: CONTEXTO E OBJETIVO ====================
    add_section_header("1. CONTEXTO E OBJETIVO ESTRATÉGICO")
    contexto_txt = data.get('contexto', '')
    if contexto_txt:
        for line in contexto_txt.split('\n'):
            line_str = line.strip()
            if line_str:
                add_p(line_str)
    else:
        add_p(f"A proposta contempla a digitalização de ponta a ponta da operação de engenharia, chão de fábrica, SGQ e frentes de canteiros da {cliente_nome}.")

    # ==================== SEÇÃO 2: ARQUITETURA E METODOLOGIA ====================
    add_section_header("2. ARQUITETURA DA SOLUÇÃO & MOBILIDADE REAL")
    add_p("O Ecossistema SINCO foi desenvolvido especificamente para sanar as lacunas operacionais entre o ERP corporativo (TOTVS Protheus) e o campo:", "Arquitetura Híbrida Inteligente: ")
    add_bullet("Aplicativo Mobile PWA com Modo Offline Nativo", "Encarregados e técnicos apontam DDS, APR, RDO e inspeções com fotos e assinaturas mesmo sem sinal de internet em canteiros remotos.")
    add_bullet("Cockpit Web Gerencial de Alta Performance", "Painel em tempo real para engenheiros coordenadores, diretores e gestores de fábrica monitorarem avanço físico e indicadores.")
    add_bullet("Barramento REST de Alta Disponibilidade", "Sincronização bidirecional, segura e auditável com as tabelas do TOTVS Protheus sem gerar lentidão nas rotinas diárias.")

    # ==================== SEÇÃO 3: ESCOPO DOS MÓDULOS FASE 1 ====================
    add_section_header("3. ESCOPO DOS MÓDULOS FASE 1 (MVP E DESENVOLVIMENTO)")
    escopo1 = data.get('escopo_fase1', [])
    for m in escopo1:
        add_p(m.get('descricao', ''), f"• {m.get('nome', 'Módulo')} ({m.get('status', 'Contemplado')}): ")
        recursos = m.get('recursos', [])
        for rec in recursos:
            if ':' in rec:
                t, d = rec.split(':', 1)
                add_bullet(f"   - {t.strip()}", d.strip())
            else:
                p_r = doc.add_paragraph()
                p_r.paragraph_format.left_indent = Inches(0.4)
                p_r.paragraph_format.space_before = Pt(1)
                p_r.paragraph_format.space_after = Pt(2)
                r_c = p_r.add_run(f"✓ {rec.strip()}")
                r_c.font.size = Pt(9)
                r_c.font.color.rgb = COLOR_TEXT

    # ==================== SEÇÃO 4: ROADMAP FUTURO (FASE 2) ====================
    add_section_header("4. ROADMAP FUTURO (FASE 2 - EXPANSÃO)")
    escopo2 = data.get('escopo_fase2', [])
    for m in escopo2:
        add_p(f"{m.get('descricao', '')} (Status: {m.get('status', 'Roadmap Futuro')})", f"• {m.get('nome', '')}: ")

    crono = data.get('investimento', {}).get('cronograma_pagamento', {}) or {}
    carencia_dias = crono.get('carencia_dias', 0)
    parcelas = crono.get('parcelas_setup', [])
    tem_entrada = any(float(p.get('dias_apos_aceite', 0)) == 0 for p in parcelas)

    # ==================== SEÇÃO 5: CONDIÇÕES COMERCIAIS & FINANCEIRAS ====================
    if tem_entrada or carencia_dias == 0:
        add_section_header("5. MODELO COMERCIAL, CONDIÇÕES DE PAGAMENTO & CRONOGRAMA")
    else:
        add_section_header(f"5. MODELO COMERCIAL, CARÊNCIA DE {carencia_dias} DIAS & CRONOGRAMA")

    # Banner de Condição Comercial
    carencia_table = doc.add_table(rows=1, cols=1)
    carencia_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    carencia_table.autofit = False
    cell_car = carencia_table.cell(0, 0)
    cell_car.width = Inches(6.8)
    set_cell_background(cell_car, "FEF3C7")  # Amber suave
    set_cell_margins(cell_car, 140, 140, 180, 180)
    p_car = cell_car.paragraphs[0]
    p_car.paragraph_format.space_after = Pt(0)

    if tem_entrada or carencia_dias == 0:
        p_ent = next((p for p in parcelas if float(p.get('dias_apos_aceite', 0)) == 0), None)
        v_ent = float(p_ent.get('valor', 0)) if p_ent else 0
        r_car_t = p_car.add_run("★ CONDIÇÃO COMERCIAL: ENTRADA FACILITADA E PARCELAMENTO EM MARCOS DE HOMOLOGAÇÃO (SEM CARÊNCIA)\n")
        r_car_t.font.bold = True
        r_car_t.font.size = Pt(10)
        r_car_t.font.color.rgb = COLOR_AMBER

        r_car_d = p_car.add_run(
            f"• Entrada Inicial: R$ {fmt_br(v_ent)} paga à vista em até {p_ent.get('dias_vencimento', 5) if p_ent else 5} dias após a assinatura formal do contrato (sem carência comercial).\n"
            f"• A LYNX e a {parceira_nome or 'consultoria parceira'} iniciam imediatamente o diagnóstico, parametrização e configuração dos módulos na fábrica.\n"
            f"• As parcelas subsequentes do setup estão atreladas aos marcos de entrega do piloto fabril (60 dias) e homologação final assistida (90 dias)."
        )
    else:
        r_car_t = p_car.add_run(f"★ CONDIÇÃO COMERCIAL DIFERENCIADA: CARÊNCIA TOTAL DE {carencia_dias} DIAS\n")
        r_car_t.font.bold = True
        r_car_t.font.size = Pt(10)
        r_car_t.font.color.rgb = COLOR_AMBER

        r_car_d = p_car.add_run(
            f"• No Dia 0 (Aceite Formal da Proposta): NENHUM desembolso financeiro é exigido da {cliente_nome} (R$ 0,00).\n"
            f"• A LYNX e a {parceira_nome or 'consultoria parceira'} iniciam imediatamente o diagnóstico, desenvolvimento e configuração dos módulos.\n"
            f"• O faturamento inicial somente tem início a partir de {carencia_dias} dias da assinatura da proposta, respeitando a entrega do MVP."
        )
    r_car_d.font.size = Pt(9)
    r_car_d.font.color.rgb = COLOR_TEXT

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(8)

    # 5.1 PARCELAS DO SETUP
    if tem_entrada or carencia_dias == 0:
        add_p("O investimento na fase de implementação é distribuído em entrada e parcelas vinculadas à evolução das entregas (sem carência comercial):", "5.1. Implementação & Setup dos Módulos: ")
    else:
        add_p(f"O investimento na fase ágil de implementação é distribuído em parcelas com carência de {carencia_dias} dias:", "5.1. Implementação & Setup do MVP: ")
    if not parcelas:
        parcelas = [
            {"nome": "Entrada", "dias_apos_aceite": 0, "dias_vencimento": 5, "valor": 9562.30},
            {"nome": "1ª Parcela Setup (60d)", "dias_apos_aceite": 60, "dias_vencimento": 10, "valor": 8520.00},
            {"nome": "2ª Parcela Setup (90d)", "dias_apos_aceite": 90, "dias_vencimento": 10, "valor": 12636.00}
        ]

    # Tabela de Parcelas Setup
    t_parc = doc.add_table(rows=1 + len(parcelas) + 1, cols=5)
    t_parc.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_parc.autofit = False
    set_table_borders(t_parc)

    headers_parc = ["Parcela / Marco", "Prazo após Aceite", "Prazo Pgto NF", "Condição", "Valor (R$)"]
    widths_parc = [Inches(2.2), Inches(1.2), Inches(1.1), Inches(1.1), Inches(1.2)]

    # Cabeçalho da tabela de parcelas
    for i, h in enumerate(headers_parc):
        c = t_parc.cell(0, i)
        c.width = widths_parc[i]
        set_cell_background(c, "F1F5F9")
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i == 4 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = COLOR_FOREST

    total_setup_calc = 0
    row_idx = 1
    for p_item in parcelas:
        nome_p = p_item.get('nome', f"Parcela {row_idx}")
        dias_a = p_item.get('dias_apos_aceite', 0)
        dias_v = p_item.get('dias_vencimento', 10)
        v = float(p_item.get('valor', 0))
        total_setup_calc += v

        row_cells = t_parc.rows[row_idx].cells
        for idx_c, w in enumerate(widths_parc):
            row_cells[idx_c].width = w
            set_cell_margins(row_cells[idx_c], 70, 70, 100, 100)

        # Col 0: Nome
        p0 = row_cells[0].paragraphs[0]
        r0 = p0.add_run(nome_p)
        r0.font.bold = True
        r0.font.size = Pt(8.5)

        # Col 1: Prazo
        p1 = row_cells[1].paragraphs[0]
        p1.add_run("No Aceite (Dia 0)" if dias_a == 0 else f"Dia {dias_a} (após {dias_a}d)").font.size = Pt(8.5)

        # Col 2: Vencimento
        p2 = row_cells[2].paragraphs[0]
        p2.add_run(f"{dias_v} dias da assinatura" if dias_a == 0 else f"{dias_v} dias da NF").font.size = Pt(8.5)

        # Col 3: Condição
        p3 = row_cells[3].paragraphs[0]
        p3.add_run("À vista (5 dias)" if dias_a == 0 else f"Fração {row_idx - 1}").font.size = Pt(8.5)

        # Col 4: Valor
        p4 = row_cells[4].paragraphs[0]
        p4.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r4 = p4.add_run(f"R$ {fmt_br(v)}")
        r4.font.bold = True
        r4.font.size = Pt(9)
        r4.font.color.rgb = COLOR_AMBER

        row_idx += 1

    # Linha Total Setup
    foot_cells = t_parc.rows[row_idx].cells
    for idx_c, w in enumerate(widths_parc):
        foot_cells[idx_c].width = w
        set_cell_background(foot_cells[idx_c], "F8FAFC")
        set_cell_margins(foot_cells[idx_c], 80, 80, 100, 100)

    p_tot_lbl = foot_cells[0].paragraphs[0]
    r_tl = p_tot_lbl.add_run(f"Total Geral Setup ({len(parcelas)} Parcelas):")
    r_tl.font.bold = True
    r_tl.font.size = Pt(8.5)

    p_tot_val = foot_cells[4].paragraphs[0]
    p_tot_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_tv = p_tot_val.add_run(f"R$ {fmt_br(total_setup_calc)}")
    r_tv.font.bold = True
    r_tv.font.size = Pt(9.5)
    r_tv.font.color.rgb = COLOR_FOREST

    # 5.2 SUSTENTAÇÃO MENSAL
    p_sp3 = doc.add_paragraph()
    p_sp3.paragraph_format.space_before = Pt(8)

    add_p("A sustentação contínua compreende suporte de nível 2 e 3, infraestrutura em nuvem dedicada, rotinas automáticas de backup e governança:", "5.2. Sustentação Mensal (SaaS & Nuvem): ")
    sust = data.get('investimento', {}).get('sustentacao_mensal', {}) or {}
    base_sm = sust.get('base_salario_minimo', sust.get('valor_base_salarios_minimos', 1.0))
    val_sm = sust.get('salario_minimo_vigente', 1518.00)
    franquia = sust.get('franquia_canteiros', 1)
    extra_sm = sust.get('canteiro_adicional_sm', sust.get('acrescimo_por_canteiro_extra_sm', 0.5))
    mens_val = float(sust.get('valor_mensal') or (base_sm * val_sm))
    extra_val = extra_sm * val_sm
    vigencia = sust.get('vigencia_meses', 24)

    add_bullet("Mensalidade de Sustentação", f"R$ {fmt_br(mens_val)}/mês (referência base de {base_sm} Salários Mínimos), cobrindo matriz, chão de fábrica e postos produtivos.")
    add_bullet("Vigência Contratual", f"Contrato com prazo de vigência de {vigencia} meses.")
    add_bullet("Critério de Reajuste", "Reajuste anual (a cada 12 meses) atrelado à variação oficial do Salário Mínimo Nacional, assegurando estabilidade sem volatilidade de índices de inflação.")
    add_bullet("Licenciamento & Acessos", "Usuários e postos touch ilimitados em toda a fábrica e escritórios, sem custos adicionais por licença.")
    add_bullet("Regra de Faturamento Mensal", f"Emissão todo dia {crono.get('mensalidade_dia_emissao', 28)} do mês corrente, com vencimento previsto para o dia {crono.get('mensalidade_dia_vencimento', 10)} do mês subsequente.")

    # 5.3 TABELA TEMPORAL DE FATURAMENTO E PAGAMENTOS (OPCIONAL)
    incluir_word = crono.get('incluir_tabela_word', True)
    if incluir_word:
        p_sp4 = doc.add_paragraph()
        p_sp4.paragraph_format.space_before = Pt(8)

        add_p("Demonstrativo cronológico do fluxo de desembolso a partir do aceite da proposta:", "5.3. Tabela Temporal de Pagamentos e Faturamento: ")

        data_base_str = crono.get('data_aceite', data_proposta)
        dt_aceite = parse_date_safe(data_base_str)

        marcos = []
        if not tem_entrada:
            marcos.append({
                "marco": "Dia 0",
                "data": dt_aceite.strftime("%d/%m/%Y"),
                "evento": "Aceite da Proposta",
                "nf": "Não aplicável (R$ 0,00)",
                "vencto": "Sem desembolso",
                "cond": "Carência Comercial Inicial",
                "valor": 0.0
            })

        for idx, p_item in enumerate(parcelas):
            d_a = p_item.get('dias_apos_aceite', 0)
            d_v = p_item.get('dias_vencimento', 10)
            val_p = float(p_item.get('valor', 0))
            nm_p = p_item.get('nome', f"Fração #{idx + 1}")

            dt_emiss = dt_aceite + timedelta(days=d_a)
            dt_venc = dt_emiss + timedelta(days=d_v)

            marcos.append({
                "marco": "No Aceite" if d_a == 0 else f"Dia {d_a}",
                "data": dt_emiss.strftime("%d/%m/%Y"),
                "evento": f"Emissão NF {nm_p}",
                "nf": dt_emiss.strftime("%d/%m/%Y"),
                "vencto": dt_venc.strftime("%d/%m/%Y"),
                "cond": f"À vista ({d_v} dias)" if d_a == 0 else f"Pagamento em {d_v} dias da NF",
                "valor": val_p
            })

        # Adiciona mensalidade a partir do 3º mês
        dt_mens1 = dt_aceite + timedelta(days=90)
        dt_venc_m1 = (dt_mens1.replace(day=1) + timedelta(days=32)).replace(day=10)
        marcos.append({
            "marco": "Mês 3+",
            "data": dt_mens1.strftime("%d/%m/%Y"),
            "evento": "Mensalidade Sustentação SaaS (24 meses)",
            "nf": f"Dia {crono.get('mensalidade_dia_emissao', 28)} de cada mês",
            "vencto": f"Dia {crono.get('mensalidade_dia_vencimento', 10)} do mês seguinte",
            "cond": "SaaS integral + usuários ilimitados",
            "valor": mens_val
        })

        # Tabela de Marcos Temporais no Word
        t_marcos = doc.add_table(rows=1 + len(marcos), cols=6)
        t_marcos.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_marcos.autofit = False
        set_table_borders(t_marcos)

        cols_m = ["Marco", "Data Estimada", "Evento / Objeto da NF", "Emissão NF", "Vencimento", "Valor Previsto"]
        widths_m = [Inches(0.9), Inches(1.1), Inches(2.0), Inches(1.0), Inches(1.0), Inches(1.1)]

        for i, h in enumerate(cols_m):
            c = t_marcos.cell(0, i)
            c.width = widths_m[i]
            set_cell_background(c, "F1F5F9")
            set_cell_margins(c, 70, 70, 80, 80)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i == 5 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(h)
            r.font.bold = True
            r.font.size = Pt(8)
            r.font.color.rgb = COLOR_FOREST

        for m_idx, m_item in enumerate(marcos):
            r_cells = t_marcos.rows[m_idx + 1].cells
            for i, w in enumerate(widths_m):
                r_cells[i].width = w
                set_cell_margins(r_cells[i], 60, 60, 80, 80)

            # Col 0: Marco
            p0 = r_cells[0].paragraphs[0]
            r0 = p0.add_run(m_item['marco'])
            r0.font.bold = True
            r0.font.size = Pt(8)

            # Col 1: Data
            r_cells[1].paragraphs[0].add_run(m_item['data']).font.size = Pt(8)

            # Col 2: Evento
            p2 = r_cells[2].paragraphs[0]
            r2_a = p2.add_run(m_item['evento'] + "\n")
            r2_a.font.bold = True
            r2_a.font.size = Pt(8)
            r2_b = p2.add_run(m_item['cond'])
            r2_b.font.size = Pt(7.5)
            r2_b.font.color.rgb = COLOR_MUTED

            # Col 3: NF
            r_cells[3].paragraphs[0].add_run(m_item['nf']).font.size = Pt(8)

            # Col 4: Vencto
            p4 = r_cells[4].paragraphs[0]
            r4 = p4.add_run(m_item['vencto'])
            r4.font.bold = True
            r4.font.size = Pt(8)

            # Col 5: Valor
            p5 = r_cells[5].paragraphs[0]
            p5.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if m_item['valor'] == 0:
                r5 = p5.add_run("R$ 0,00")
                r5.font.bold = True
                r5.font.size = Pt(8)
                r5.font.color.rgb = COLOR_MUTED
            else:
                r5 = p5.add_run(f"R$ {fmt_br(m_item['valor'])}")
                r5.font.bold = True
                r5.font.size = Pt(8.5)
                r5.font.color.rgb = COLOR_FOREST

    # ==================== SEÇÃO 6: LICENCIAMENTO, SLA E LGPD ====================
    add_section_header("6. LICENCIAMENTO, SEGURANÇA, SLA E LGPD")
    add_bullet("Licenciamento Empresarial Sem Trava por Usuário", f"Sem limite de colaboradores cadastrados na {cliente_nome}. Todos os operadores, engenheiros e encarregados podem acessar a plataforma.")
    add_bullet("SLA de Suporte e Atendimento Técnico", "Suporte prioritário de segunda a sexta das 08h às 18h com atendimento via canal direto e monitoramento proativo de barramentos.")
    add_bullet("Segurança, Criptografia e Backups", "Banco de dados em nuvem com TLS 1.3, backups diários automatizados e retenção histórica para proteção de passivos contratuais e auditorias.")
    add_bullet("Conformidade Total com a LGPD", "Dados pessoais de operadores e encarregados coletados estritamente para finalidades legais e de SST, em conformidade com a Lei 13.709/2018.")

    # ==================== SEÇÃO 7: DE ACORDO E ASSINATURAS ====================
    add_section_header("7. TERMO DE ACEITE E DE ACORDO DAS PARTES")
    if tem_entrada or carencia_dias == 0:
        add_p(f"Estando as partes de pleno acordo com o escopo técnico, metodologia de implantação, condições de pagamento acordadas (sem carência comercial) e governança compartilhada, assinam o presente documento:")
    else:
        add_p(f"Estando as partes de pleno acordo com o escopo técnico, metodologia ágil, carência de {carencia_dias} dias, condições comerciais e governança compartilhada, assinam o presente documento:")

    p_sp5 = doc.add_paragraph()
    p_sp5.paragraph_format.space_before = Pt(30)

    # Assinaturas em tabela de 3 colunas (Lynx EMS, SIM Consultoria, Cliente)
    num_cols = 3 if parceira_habilitada else 2
    sig_table = doc.add_table(rows=1, cols=num_cols)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False

    w_col = Inches(6.8 / num_cols)

    # Col 1: LYNX EMS
    c1 = sig_table.cell(0, 0)
    c1.width = w_col
    set_cell_margins(c1, 40, 40, 40, 40)
    p_sig1 = c1.paragraphs[0]
    p_sig1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sig1.add_run("_____________________________________\n").font.color.rgb = COLOR_MUTED
    r_s1_t = p_sig1.add_run(f"{fornecedor_nome}\n")
    r_s1_t.font.bold = True
    r_s1_t.font.size = Pt(8.5)
    r_s1_t.font.color.rgb = COLOR_PRIMARY
    p_sig1.add_run("Tecnologia & Desenvolvimento de Software").font.size = Pt(7.5)

    if parceira_habilitada:
        # Col 2: SIM CONSULTORIA
        c2 = sig_table.cell(0, 1)
        c2.width = w_col
        set_cell_margins(c2, 40, 40, 40, 40)
        p_sig2 = c2.paragraphs[0]
        p_sig2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sig2.add_run("_____________________________________\n").font.color.rgb = COLOR_MUTED
        r_s2_t = p_sig2.add_run(f"{parceira_nome}\n")
        r_s2_t.font.bold = True
        r_s2_t.font.size = Pt(8.5)
        r_s2_t.font.color.rgb = COLOR_PRIMARY
        p_sig2.add_run("Consultoria Protheus & Implantação").font.size = Pt(7.5)

        # Col 3: CLIENTE
        c3 = sig_table.cell(0, 2)
        c3.width = w_col
        set_cell_margins(c3, 40, 40, 40, 40)
        p_sig3 = c3.paragraphs[0]
        p_sig3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sig3.add_run("_____________________________________\n").font.color.rgb = COLOR_MUTED
        r_s3_t = p_sig3.add_run(f"{cliente_nome}\n")
        r_s3_t.font.bold = True
        r_s3_t.font.size = Pt(8.5)
        r_s3_t.font.color.rgb = COLOR_PRIMARY
        p_sig3.add_run(f"Diretoria / Contratante").font.size = Pt(7.5)
    else:
        # Col 2: CLIENTE
        c2 = sig_table.cell(0, 1)
        c2.width = w_col
        set_cell_margins(c2, 40, 40, 40, 40)
        p_sig2 = c2.paragraphs[0]
        p_sig2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sig2.add_run("_____________________________________\n").font.color.rgb = COLOR_MUTED
        r_s2_t = p_sig2.add_run(f"{cliente_nome}\n")
        r_s2_t.font.bold = True
        r_s2_t.font.size = Pt(8.5)
        r_s2_t.font.color.rgb = COLOR_PRIMARY
        p_sig2.add_run("Diretoria / Contratante").font.size = Pt(7.5)

    doc.save(output_path)
    print(f"Documento Word gerado com sucesso em: {output_path}")
    return output_path

if __name__ == "__main__":
    json_arg = sys.argv[1] if len(sys.argv) > 1 else None
    out_arg = sys.argv[2] if len(sys.argv) > 2 else None
    generate_proposal_docx(json_arg, out_arg)
